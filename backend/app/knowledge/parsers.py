"""Archivos a texto para la base de conocimiento. Sin dependencias nuevas: pypdf y openpyxl ya están en el
proyecto; Word (.docx) se lee directamente del XML del paquete."""

import csv
import html
import io
import re
import zipfile
from xml.etree import ElementTree

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_SHEET_ROWS = 5000


class ParseError(Exception):
    """El archivo no se puede leer (mensaje para el usuario)."""


SUPPORTED = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".json": "application/json",
}


def extension(filename: str) -> str:
    m = re.search(r"\.[A-Za-z0-9]+$", filename or "")
    return m.group(0).lower() if m else ""


def html_to_text(page: str) -> str:
    """HTML → texto conservando títulos como Markdown (sirven de secciones al fragmentar)."""
    page = re.sub(r"(?is)<(script|style|noscript|svg|template|nav|footer)[^>]*>.*?</\1>", " ", page or "")
    page = re.sub(r"(?is)<h([1-6])[^>]*>(.*?)</h\1>",
                  lambda m: "\n\n" + "#" * int(m.group(1)) + " " + re.sub(r"(?s)<[^>]+>", " ", m.group(2)) + "\n\n",
                  page)
    page = re.sub(r"(?is)<li[^>]*>", "\n- ", page)
    page = re.sub(r"(?is)<(br|p|div|tr|section|article|table|ul|ol)[^>]*>", "\n\n", page)
    page = re.sub(r"(?is)</(td|th)>", " | ", page)
    page = re.sub(r"(?s)<[^>]+>", " ", page)
    page = html.unescape(page)
    lines = [re.sub(r"[ \t\r\f\v]+", " ", line).strip() for line in page.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def html_title(page: str) -> str | None:
    m = re.search(r"(?is)<title[^>]*>(.*?)</title>", page or "")
    if not m:
        return None
    return html.unescape(m.group(1)).strip() or None


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "").strip() for p in reader.pages]
    except Exception as e:  # noqa: BLE001 (pypdf lanza muchos tipos)
        raise ParseError(f"No se pudo leer el PDF: {e}") from e
    text = "\n\n".join(t for t in pages if t)
    if not text.strip():
        raise ParseError("El PDF no tiene texto extraíble (¿es un escaneo?)")
    return text


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            root = ElementTree.fromstring(z.read("word/document.xml"))
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError) as e:
        raise ParseError("No se pudo leer el documento de Word (.docx)") from e
    out: list[str] = []
    for p in root.iter(f"{_W}p"):
        style = p.find(f"{_W}pPr/{_W}pStyle")
        text = "".join(t.text or "" for t in p.iter(f"{_W}t")).strip()
        if not text:
            continue
        level = None
        if style is not None:
            m = re.search(r"(?i)heading\s*(\d)|t[ií]tulo\s*(\d)", style.get(f"{_W}val", ""))
            if m:
                level = int(m.group(1) or m.group(2))
        out.append(("#" * min(level, 6) + " " + text) if level else text)
    return "\n\n".join(out)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ParseError(f"No se pudo leer el Excel: {e}") from e
    blocks: list[str] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True, max_row=MAX_SHEET_ROWS))
        rows = [r for r in rows if any(c not in (None, "") for c in r)]
        if not rows:
            continue
        blocks.append(f"# {ws.title}\n\n" + _table(rows))
    return "\n\n".join(blocks)


def _table(rows: list) -> str:
    """Filas → «Encabezado: valor; …» (una línea por fila: cada fila queda autocontenida al fragmentar)."""
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    has_header = all(h and not re.fullmatch(r"[\d.,$ ]+", h) for h in header if h) and any(header)
    lines = []
    for r in rows[1:] if has_header else rows:
        cells = ["" if c is None else str(c).strip() for c in r]
        if has_header:
            pairs = [f"{h}: {v}" for h, v in zip(header, cells, strict=False) if v]
            lines.append("; ".join(pairs))
        else:
            lines.append(" | ".join(c for c in cells if c))
    return "\n".join(line for line in lines if line)


def _csv(data: bytes) -> str:
    text = data.decode("utf-8-sig", "replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))[:MAX_SHEET_ROWS]
    rows = [r for r in rows if any(c.strip() for c in r)]
    return _table(rows) if rows else ""


def parse(filename: str, data: bytes) -> tuple[str, str]:
    """(texto, mime). Lanza ParseError con un mensaje claro si el formato no se soporta o no tiene texto."""
    if len(data) > MAX_FILE_BYTES:
        raise ParseError(f"El archivo supera {MAX_FILE_BYTES // (1024 * 1024)} MB")
    ext = extension(filename)
    if ext not in SUPPORTED:
        raise ParseError("Formato no soportado. Usa PDF, Word (.docx), Excel (.xlsx), CSV, TXT, Markdown o HTML")
    if ext == ".pdf":
        text = _pdf(data)
    elif ext == ".docx":
        text = _docx(data)
    elif ext == ".xlsx":
        text = _xlsx(data)
    elif ext == ".csv":
        text = _csv(data)
    elif ext in (".html", ".htm"):
        text = html_to_text(data.decode("utf-8", "replace"))
    else:
        text = data.decode("utf-8-sig", "replace")
    if not text.strip():
        raise ParseError("El archivo no tiene texto")
    return text, SUPPORTED[ext]
