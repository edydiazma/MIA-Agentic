"""Catálogo de productos: importación (CSV/Excel/JSON, catálogo de Meta, feed externo), búsqueda y descripción.

La búsqueda usa el índice de texto completo en español (products.search, columna generada) y, si no hay
coincidencias, similitud por trigramas sobre el nombre.
"""

import asyncio
import csv
import io
import json
import logging
import unicodedata

import httpx
from sqlalchemy import func, literal_column, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db import SessionLocal
from app.fields import parse_number
from app.models import CatalogSyncRun, Organization, Product, utcnow
from app.settings_store import get_setting, set_setting

log = logging.getLogger(__name__)

# Encabezados aceptados (en minúscula, sin tildes) -> atributo del producto
ALIASES = {
    "sku": "sku", "codigo": "sku", "referencia": "sku", "ref": "sku", "id": "sku", "retailer_id": "sku",
    "nombre": "name", "name": "name", "titulo": "name", "title": "name", "producto": "name",
    "descripcion": "description", "description": "description",
    "categoria": "category", "category": "category", "linea": "category",
    "marca": "brand", "brand": "brand",
    "precio": "price", "price": "price", "valor": "price",
    "precio_oferta": "sale_price", "sale_price": "sale_price", "oferta": "sale_price",
    "moneda": "currency", "currency": "currency",
    "stock": "stock", "inventario": "stock", "cantidad": "stock", "existencias": "stock",
    "disponible": "available", "availability": "available", "disponibilidad": "available",
    "url": "url", "link": "url", "enlace": "url",
    "imagen": "image_url", "image_url": "image_url", "image_link": "image_url", "foto": "image_url",
}
UNAVAILABLE = {"no", "false", "0", "agotado", "out of stock", "out_of_stock", "discontinued", "no disponible"}
SEARCH = literal_column("products.search")  # tsvector generado (no mapeado en el ORM)


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s).strip().lower())
    return "".join(c for c in s if not unicodedata.combining(c)).replace(" ", "_")


def map_row(row: dict, default_currency: str) -> dict | None:
    """Convierte una fila con encabezados libres en los campos de Product (+ atributos extra)."""
    data: dict = {"attributes": {}}
    for header, value in row.items():
        if header is None or value is None or str(value).strip() == "":
            continue
        key = ALIASES.get(_norm(header))
        if key:
            data[key] = value
        else:
            data["attributes"][str(header).strip()] = str(value).strip()
    if not data.get("name"):
        return None
    data["sku"] = str(data.get("sku") or _norm(data["name"]))[:100]
    data["name"] = str(data["name"]).strip()[:255]
    for k in ("price", "sale_price"):
        if k in data:
            n = parse_number(str(data[k]))
            data[k] = round(n, 2) if n is not None and n >= 0 else None
    if "stock" in data:
        n = parse_number(str(data["stock"]))
        data["stock"] = int(n) if n is not None else None
    if "available" in data:
        data["available"] = _norm(data["available"]).replace("_", " ") not in UNAVAILABLE
    elif data.get("stock") is not None:
        data["available"] = data["stock"] > 0
    data["currency"] = str(data.get("currency") or default_currency).upper()[:3]
    for k in ("description", "category", "brand", "url", "image_url"):
        if k in data:
            data[k] = str(data[k]).strip() or None
    data["attributes"] = data["attributes"] or None
    return data


def parse_file(filename: str, content: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h or "").strip() for h in rows[0]]
        return [dict(zip(headers, r, strict=False)) for r in rows[1:] if any(c not in (None, "") for c in r)]
    if name.endswith(".json"):
        data = json.loads(content)
        return data if isinstance(data, list) else data.get("products") or data.get("data") or []
    text = content.decode("utf-8-sig", "replace")
    if not text.strip():
        return []
    dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


async def upsert(session: AsyncSession, org: int, rows: list[dict], source: str, replace_source: bool = False,
                 agent_id: int | None = None) -> dict:
    """Crea o actualiza por SKU y registra la ejecución en catalog_sync_runs.

    Con replace_source, los productos de esa misma fuente que no vinieron quedan no disponibles.
    """
    run = CatalogSyncRun(organization_id=org, source="import" if source == "import" else source, created_by=agent_id)
    session.add(run)
    await session.flush()
    cfg = await get_setting(session, "catalog", org)
    existing = {p.sku: p for p in (await session.scalars(select(Product).where(Product.organization_id == org))).all()}
    created = updated = invalid = 0
    seen: set[str] = set()
    for row in rows:
        data = map_row(row, cfg["currency"])
        if not data:
            invalid += 1
            continue
        seen.add(data["sku"])
        product = existing.get(data["sku"])
        if not product:
            product = Product(organization_id=org, source=source)
            session.add(product)
            existing[data["sku"]] = product
            created += 1
        else:
            updated += 1
        for k, v in data.items():
            setattr(product, k, v)
        if source == "meta":
            product.in_meta_catalog = True
    disabled = 0
    if replace_source:
        for p in existing.values():
            if p.source == source and p.sku not in seen and p.available:
                p.available = False
                disabled += 1
    stats = {"created": created, "updated": updated, "invalid": invalid, "disabled": disabled}
    run.status, run.stats, run.finished_at = "done", stats, utcnow()
    await session.commit()
    return {**stats, "run_id": run.id}


async def _failed_run(session: AsyncSession, org: int, source: str, error: str, agent_id: int | None = None) -> None:
    session.add(CatalogSyncRun(organization_id=org, source=source, status="failed", error=error[:2000],
                               created_by=agent_id, finished_at=utcnow()))
    await session.commit()


# --- Búsqueda (la usan los agentes de IA y los flujos) ---------------------------
async def search(session: AsyncSession, org: int, query: str, max_price: float | None = None,
                 category: str | None = None, limit: int = 5) -> list[Product]:
    base = select(Product).where(Product.organization_id == org, Product.available)
    if max_price:
        base = base.where(func.coalesce(Product.sale_price, Product.price) <= max_price)
    if category:
        base = base.where(Product.category.ilike(f"%{category.strip()}%"))
    q = (query or "").strip()
    if not q:
        return list((await session.scalars(base.order_by(Product.name).limit(limit))).all())

    tsq = func.websearch_to_tsquery("spanish", q)
    rows = (await session.scalars(
        base.where(SEARCH.op("@@")(tsq)).order_by(func.ts_rank(SEARCH, tsq).desc(), Product.name).limit(limit))).all()
    if rows:
        return list(rows)
    # Sin coincidencias de texto completo: similitud aproximada (errores de tipeo, modelos parciales)
    sim = func.extensions.similarity(Product.name, q)
    rows = (await session.scalars(
        base.where(or_(Product.name.ilike(f"%{q}%"), Product.sku.ilike(f"%{q}%"), sim > 0.2))
        .order_by(sim.desc(), Product.name).limit(limit))).all()
    return list(rows)


def describe(p: Product) -> str:
    price = p.sale_price or p.price
    parts = [f"[{p.sku}] {p.name}"]
    if price is not None:
        parts.append(f"{p.currency} {float(price):,.0f}".replace(",", ".") + (" (oferta)" if p.sale_price else ""))
    if p.stock is not None:
        parts.append(f"stock {p.stock}")
    if p.category:
        parts.append(p.category)
    if p.attributes:
        parts.append(", ".join(f"{k}: {v}" for k, v in list(p.attributes.items())[:8]))
    if p.description:
        parts.append(p.description[:300])
    if p.url:
        parts.append(p.url)
    return " | ".join(parts)


# --- Sincronizaciones -----------------------------------------------------------
async def sync_meta(session: AsyncSession, org: int, client, agent_id: int | None = None) -> dict:
    """Trae los productos del catálogo de Commerce Manager conectado a WhatsApp."""
    cfg = await get_setting(session, "catalog", org)
    if not cfg["meta_catalog_id"]:
        raise ValueError("Configura el ID del catálogo de Meta")
    try:
        items = await client.list_catalog_products(cfg["meta_catalog_id"])
    except Exception as e:
        await _failed_run(session, org, "meta", f"{type(e).__name__}: {e}", agent_id)
        raise
    rows = [{
        "retailer_id": i.get("retailer_id"), "name": i.get("name"), "description": i.get("description"),
        "price": i.get("price"), "sale_price": i.get("sale_price"), "currency": i.get("currency"),
        "availability": i.get("availability"), "url": i.get("url"), "image_url": i.get("image_url"),
        "brand": i.get("brand"), "category": i.get("category"),
    } for i in items]
    result = await upsert(session, org, rows, "meta", replace_source=True, agent_id=agent_id)
    await set_setting(session, "catalog", {"last_sync": {"source": "meta", "at": utcnow().isoformat(), **result}},
                      org=org, agent_id=agent_id, source="system")
    return result


async def sync_feed(session: AsyncSession, org: int, agent_id: int | None = None) -> dict:
    """Sistema externo (ERP, Shopify, inventario): descarga un CSV o JSON publicado en una URL."""
    cfg = await get_setting(session, "catalog", org)
    if not str(cfg["feed_url"]).startswith("http"):
        raise ValueError("Configura la URL del feed (https://...)")
    try:
        async with httpx.AsyncClient(timeout=60, follow_redirects=True) as http:
            r = await http.get(cfg["feed_url"])
            r.raise_for_status()
        rows = parse_file("feed." + ("json" if cfg["feed_format"] == "json" else "csv"), r.content)
    except Exception as e:
        await _failed_run(session, org, "feed", f"{type(e).__name__}: {e}", agent_id)
        raise
    result = await upsert(session, org, rows, "feed", replace_source=True, agent_id=agent_id)
    await set_setting(session, "catalog", {"last_sync": {"source": "feed", "at": utcnow().isoformat(), **result}},
                      org=org, agent_id=agent_id, source="system")
    return result


async def feed_loop() -> None:
    """Sincroniza el feed externo de cada organización cada `feed_interval_hours`."""
    last: dict[int, float] = {}
    while True:
        await asyncio.sleep(300)
        try:
            async with SessionLocal() as session:
                orgs = (await session.scalars(select(Organization.id))).all()
                for org in orgs:
                    cfg = await get_setting(session, "catalog", org)
                    hours = float(cfg.get("feed_interval_hours") or 0)
                    now = asyncio.get_running_loop().time()
                    if hours > 0 and cfg["feed_url"] and now - last.get(org, 0) >= hours * 3600:
                        last[org] = now
                        try:
                            log.info("Catálogo org %s: %s", org, await sync_feed(session, org))
                        except Exception:
                            log.exception("Falló la sincronización del feed (org %s)", org)
        except Exception:
            log.exception("Falló el ciclo de sincronización del catálogo")
